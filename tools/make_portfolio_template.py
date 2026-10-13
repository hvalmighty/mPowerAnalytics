"""Builds public/sample/Portfolio_Template.xlsx: the minimum portfolio input, using the
PMS security codes found in the market price export (SMM = SM Investments, ACY = Ayala Corp,
SMP = SM Prime, ARE = AREIT, CNV = Converge). Quantities are illustrative.
Run: python3 tools/make_portfolio_template.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "public" / "sample" / "Portfolio_Template.xlsx"
F = "Arial"
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
REQ = PatternFill("solid", fgColor="0D366B")
OPT = PatternFill("solid", fgColor="1C5CAB")
INPUT = Font(name=F, size=10, color="0000FF")
NOTE = Font(name=F, size=9, italic=True, color="52514E")
B = Border(bottom=Side(style="thin", color="E2E0DA"))

wb = Workbook()
ws = wb.active
ws.title = "Holdings"
cols = [
    ("Security", 12, True, "PMS security code, exactly as in the market price file (e.g. SMM, ACY, BDO). Cash = CASH"),
    ("Security_Name", 40, False, "Optional, shown in the reports"),
    ("Sector", 16, False, "Recommended: Financials, Industrial, Holding Firms, Property, Services, Mining and Oil"),
    ("Quantity", 14, True, "Shares held; for CASH the PHP balance"),
    ("Avg_Cost", 12, False, "Optional"),
    ("Market_Value", 16, False, "Optional: PMS market value, cross-checked against the price file"),
]
for j, (h, w, req, note) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.font, c.fill = HDR, (REQ if req else OPT)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.comment = Comment(note, "Template")
    ws.column_dimensions[c.column_letter].width = w
ws.freeze_panes = "A2"
rows = [
    ("BDO", "BDO Unibank, Inc.", "Financials", 60000, 128.40),
    ("BPI", "Bank of the Philippine Islands", "Financials", 50000, 112.10),
    ("MBT", "Metropolitan Bank & Trust Co.", "Financials", 80000, 64.25),
    ("CBC", "China Banking Corporation", "Financials", 40000, 61.00),
    ("ICT", "International Container Terminal Services, Inc.", "Services", 20000, 385.00),
    ("TEL", "PLDT Inc.", "Services", 2500, 1325.00),
    ("GLO", "Globe Telecom, Inc.", "Services", 1500, 1800.00),
    ("CNV", "Converge ICT Solutions, Inc.", "Services", 150000, 14.10),
    ("SMM", "SM Investments Corporation", "Holding Firms", 6000, 845.50),
    ("ACY", "Ayala Corporation", "Holding Firms", 7000, 610.00),
    ("JGS", "JG Summit Holdings, Inc.", "Holding Firms", 100000, 27.80),
    ("ALI", "Ayala Land, Inc.", "Property", 250000, 28.90),
    ("SMP", "SM Prime Holdings, Inc.", "Property", 300000, 24.60),
    ("ARE", "AREIT, Inc.", "Property", 120000, 35.20),
    ("JFC", "Jollibee Foods Corporation", "Industrial", 20000, 215.00),
    ("URC", "Universal Robina Corporation", "Industrial", 40000, 92.40),
    ("MER", "Manila Electric Company", "Industrial", 8000, 505.00),
    ("SCC", "Semirara Mining and Power Corporation", "Mining and Oil", 60000, 31.50),
    ("CASH", "Cash (PHP)", "Cash", 4000000, None),
]
for i, r in enumerate(rows, 2):
    for j, v in enumerate(r, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font, c.border = INPUT, B
    ws.cell(row=i, column=4).number_format = "#,##0"
    ws.cell(row=i, column=5).number_format = "#,##0.00"
dv = DataValidation(type="list", formula1='"Financials,Industrial,Holding Firms,Property,Services,Mining and Oil,Cash"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("C2:C2000")
n = len(rows) + 3
for k, t in enumerate([
    "Only Security and Quantity are required. Quantities and costs in this template are illustrative.",
    "Security codes must match the market price file (the PMS codes, e.g. SMM for SM Investments, ACY for Ayala Corp, SMP for SM Prime).",
    "Optional extra sheets, same layout as the full sample workbook: Transactions, Cash_Flows, Benchmark_Constituents, Portfolio_Info.",
]):
    ws.cell(row=n + k, column=1, value=t).font = NOTE
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("wrote", OUT)
